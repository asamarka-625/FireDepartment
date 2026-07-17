# Внешние зависимости
from typing import Literal, Any, List
from datetime import datetime
from zoneinfo import ZoneInfo
from io import BytesIO
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
# Внутренние модули
from web_app.src.core import cfg
from web_app.src.schemas import ReportScheme
from web_app.src.utils.sorting import sort_special_first, has_adjacent_uppercase


MONTHS_RU = {
    1: "января", 2: "февраля", 3: "марта", 4: "апреля",
    5: "мая", 6: "июня", 7: "июля", 8: "августа",
    9: "сентября", 10: "октября", 11: "ноября", 12: "декабря"
}


class ReportExelCreator:
    def __init__(self):
        self.wb = load_workbook(f"{cfg.TEMPLATE_NAME}.xlsx")
        self.ws = self.wb.active

        self.font = Font(name="Times New Roman", size=12, bold=False, italic=False)
        self.title_font = Font(name="Times New Roman", size=12, bold=True, italic=False)

        thin = Side(style="thin")
        self.border = Border(left=thin, right=thin, top=thin, bottom=thin)

        self.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.number_format = "General"

        self.fills = {
            "none": PatternFill(fill_type=None),
            "green": PatternFill(start_color="C8E7A8", end_color="C8E7A8", fill_type="solid"),
            "red": PatternFill(start_color="FF8080", end_color="FF8080", fill_type="solid"),
        }

    def create_cell(
        self,
        coordinate: str,
        value: Any,
        color_fill: Literal["none", "green", "red"] = "none",
        title_font: bool = False
    ):
        self.ws[coordinate].value = value
        self.ws[coordinate].font = self.title_font if title_font else self.font
        self.ws[coordinate].fill = self.fills.get(color_fill, PatternFill(fill_type=None))
        self.ws[coordinate].border = self.border
        self.ws[coordinate].alignment = self.alignment
        self.ws[coordinate].number_format = self.number_format

    def create_merged_cell(
        self,
        cell_range: str,
        value: Any,
        color_fill: Literal["none", "green", "red"] = "none",
        title_font: bool = False
    ):
        self.ws.merge_cells(cell_range)

        font = self.title_font if title_font else self.font
        fill = self.fills.get(color_fill, PatternFill(fill_type=None))

        for row in self.ws[cell_range]:
            for cell in row:
                cell.font = font
                cell.fill = fill
                cell.border = self.border
                cell.alignment = self.alignment
                cell.number_format = self.number_format

        top_left = cell_range.split(":")[0]
        self.ws[top_left].value = value

    # Создаем строку с данными ПСЧ. Возвращает номер последней занятой строки.
    def create_section(
        self,
        section_num: int,
        start_row: int,
        report: ReportScheme
    ) -> int:
        machinery_sorted = sort_special_first(report.machinery, lambda m: m.title)

        machines = [
            {
                "title": m.title,
                "status": m.status,
                "current_personnel": m.current_personnel,
                "gdzs": m.gdzs,
                "supervisor": m.supervisor
            }
            for m in machinery_sorted
        ]

        len_row = len(machines)
        end_row = start_row + len_row - 1

        inactive_percent = round(
            ((report.personnel - report.current_personnel) * 100) / report.personnel, 2
        ) if report.personnel > 0 else 100

        active_personal_machines = sum(
            m["current_personnel"] for m in machines if m["status"].lower() == "включено"
        )

        operational_value = "Да" if report.operational_machinery else "Нет"

        # --- Общие для отделения ячейки (объединяем по вертикали на все машины) ---
        section_cells = [
            ("A", "A", section_num, False),
            ("B", "C", report.section.upper(), True),
            ("D", "D", report.total_personnel, False),    # По штату
            ("E", "E", report.personnel, False),          # По списку
            ("F", "F", report.current_personnel, False),  # На лицо
            ("G", "G", f"{inactive_percent}%", False),    # Отсутствуют
            ("M", "M", active_personal_machines, False),  # Всего
            ("N", "O", operational_value, False),         # Оперативная машина (общая)
            ("P", "Q", report.leadership, False),         # Руководство караула
        ]

        for col_start, col_end, value, is_title in section_cells:
            self.create_merged_cell(
                cell_range=f"{col_start}{start_row}:{col_end}{end_row}",
                value=value,
                title_font=is_title
            )

        # --- Данные по каждой машине ---
        for i, machine in enumerate(machines):
            row = start_row + i

            self.create_cell(f"H{row}", machine["title"])   # Техника

            status_lower = machine["status"].lower()
            personnel_color: Literal["none", "green", "red"] = "none"

            if status_lower == "включено":
                personnel_value = machine["current_personnel"]
            else:
                personnel_value = machine["status"].upper()

                if status_lower in ("резерв", "резерв (лсо)"):
                    personnel_color = "green"
                else:
                    personnel_color = "red"

            self.create_cell(f"I{row}", personnel_value, color_fill=personnel_color)  # Личный состав
            self.create_cell(f"J{row}", machine["gdzs"])                              # ГДЗС

            self.create_merged_cell(
                cell_range=f"K{row}:L{row}",
                value=machine["supervisor"]
            ) # Старший на машине

        return end_row

    # Возвращает готовый xlsx в памяти, не сохраняя на диск
    def run(
        self,
        reports: List[ReportScheme],
        machineries: List[dict],
        creator: str
    ) -> BytesIO:
        # свежая книга на каждый экспорт (иначе singleton копит листы/строки)
        self.wb = load_workbook(f"{cfg.TEMPLATE_NAME}.xlsx")

        # активный лист шаблона -> строевая записка (2-й лист)
        report_ws = self.wb.active
        report_ws.title = "Строевая записка"
        self.ws = report_ws
        self._fill_report(reports, creator)

        # новый лист с техникой -> ставим первым
        machines_ws = self.wb.create_sheet("Техника", 0)
        self.ws = machines_ws
        self._fill_machineries(machineries)

        # файл открывается на листе с техникой
        self.wb.active = 0

        buffer = BytesIO()
        self.wb.save(buffer)
        buffer.seek(0)
        return buffer

    def _fill_report(
        self,
        reports: List[ReportScheme],
        creator: str
    ) -> None:
        now = datetime.now(ZoneInfo("Europe/Moscow"))

        hours = now.strftime('%H')
        minutes = now.strftime('%M')

        day = now.strftime('%d')
        month = MONTHS_RU[now.month]
        year = now.strftime('%Y')

        self.create_cell("A2", f"На {hours} часов {minutes} минут", title_font=True)
        self.create_cell("E2", day, title_font=True)
        self.create_cell("H2", month, title_font=True)
        self.create_cell("K2", year, title_font=True)

        start_row = 6  # первая строка под данные (после шапки)

        for i, report in enumerate(reports):
            end_row = self.create_section(
                section_num=i + 1,
                start_row=start_row,
                report=report
            )
            start_row = end_row + 1  # следующая секция — сразу под текущей

        self.create_merged_cell(
            cell_range=f"A{start_row}:B{start_row}",
            value="Строевую записку подготовил"
        )
        self.create_merged_cell(
            cell_range=f"C{start_row}:N{start_row}",
            value=""
        )
        self.create_merged_cell(
            cell_range=f"O{start_row}:Q{start_row}",
            value=creator
        )

    # Вся техника с обслуживанием (первый лист)
    def _fill_machineries(
        self,
        machineries: List[dict]
    ) -> None:
        headers = [
            "Подразделение", "Название", "Модель", "Номер", "Статус",
            "Обслуживание", "Дата обслуживания"
        ]
        cols = ["A", "B", "C", "D", "E", "F", "G"]

        widths = {
            "A": 24, "B": 26, "C": 22, "D": 16,
            "E": 16, "F": 26, "G": 16
        }
        for col, width in widths.items():
            self.ws.column_dimensions[col].width = width

        for col, title in zip(cols, headers):
            self.create_cell(f"{col}1", title, title_font=True)

        self.ws.freeze_panes = "A2"

        if not machineries:
            return

        machineries = sorted(
            machineries,
            key=lambda m: (m["section"], 0 if has_adjacent_uppercase(m["title"]) else 1)
        )

        start_row = 2

        for idx, m in enumerate(machineries):
            row = start_row + idx

            # цвет статуса
            status_lower = m["status"].lower()
            if status_lower == "включено":
                status_color: Literal["none", "green", "red"] = "none"
            elif status_lower in ("резерв", "резерв (лсо)"):
                status_color = "green"
            else:
                status_color = "red"

            self.create_cell(f"B{row}", m["title"])
            self.create_cell(f"C{row}", m["model"] or "")
            self.create_cell(f"D{row}", m["number"] or "")
            self.create_cell(f"E{row}", m["status"], color_fill=status_color)
            self.create_cell(f"F{row}", m["maintenance_note"] or "")

            date_val = m["maintenance_date"]
            self.create_cell(
                f"G{row}",
                date_val.strftime("%d.%m.%Y") if date_val else ""
            )

        # объединяем "Подразделение" по группам подряд идущих машин
        group_start = start_row
        total = len(machineries)

        for idx in range(total):
            row = start_row + idx
            is_last = idx == total - 1
            next_section = None if is_last else machineries[idx + 1]["section"]

            if is_last or machineries[idx]["section"] != next_section:
                self.create_merged_cell(
                    f"A{group_start}:A{row}",
                    machineries[idx]["section"],
                    title_font=True
                )
                group_start = row + 1